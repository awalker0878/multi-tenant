"""Build only the v1.1 developed documents; never mutate the frozen reference or workbooks."""
from pathlib import Path
import json, os, itertools
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT=Path(__file__).resolve().parents[1]
BOOKMARK_IDS=itertools.count(100); FIELD_IDS=itertools.count(10000)
NAVY='15384A';TEAL='0B6B70';INK='21333C';MUTED='52636D'
BASEMAP={
'RA':('00_Reference_Architecture_v1_4.docx','RA_s_{:03d}'),
'GM':('01_Gap_Map_and_Decision_Register_v1_4.docx','GM_s_{:03d}'),
'NET':('02_Fabric_Security_and_Interfaces_v1_4.docx','NET_s_{:03d}'),
'VND':('03_Vendor_Stack_Realizations_v1_4.docx','VND_s_{:03d}'),
'PROV':('04_Provisioning_and_Commissioning_v1_4.docx','PROV_s_{:03d}'),
'SVC':('05_Shared_Services_Data_and_Recovery_v1_4.docx','SVC_s_{:03d}'),
'QUAL':('06_Site_Design_Qualification_and_Operations_v1_4.docx','QUAL_s_{:03d}'),
'WD':('07_Worked_Infrastructure_Design_and_Acceptance_v1_4.docx','WD14_S{:02d}')}
OLD={
'AK':'01_Architecture/Architecture_Kit.docx','AT':'01_Architecture/HLD_and_Architecture_Review_Template.docx',
'EK':'02_Engineering/Engineering_Kit.docx','ET':'02_Engineering/LLD_and_Engineering_Review_Template.docx',
'VC':'02_Engineering/Vendor_Realization_Cards.docx','IK':'03_Implementation/Implementation_Kit.docx',
'IT':'03_Implementation/MOP_Test_and_Handover_Template.docx','EX':'04_Shared/Worked_Delivery_Example.docx'}

def el(tag, **attrs):
    e=OxmlElement(tag)
    for k,v in attrs.items(): e.set(qn(k),str(v))
    return e

def bookmark(p,name):
    i=next(BOOKMARK_IDS)
    p._p.append(el('w:bookmarkStart',**{'w:id':i,'w:name':name}));p._p.append(el('w:bookmarkEnd',**{'w:id':i}))

def link(p,label,target=None,anchor=None):
    h=el('w:hyperlink')
    if target: h.set(qn('r:id'),p.part.relate_to(target+('#'+anchor if anchor else ''),RT.HYPERLINK,is_external=True))
    elif anchor:h.set(qn('w:anchor'),anchor)
    r=el('w:r');pr=el('w:rPr');pr.append(el('w:color',**{'w:val':TEAL}));pr.append(el('w:u',**{'w:val':'single'}));r.append(pr)
    t=el('w:t');t.text=label;r.append(t);h.append(r);p._p.append(h)

def response(p,tag):
    s=el('w:sdt');pr=el('w:sdtPr')
    for name,val in [('w:id',next(FIELD_IDS)),('w:tag',tag),('w:alias',tag.replace('_',' '))]:pr.append(el(name,**{'w:val':val}))
    pr.append(el('w:text',**{'w:multiLine':'1'}));s.append(pr)
    c=el('w:sdtContent');r=el('w:r');rp=el('w:rPr');rp.append(el('w:color',**{'w:val':TEAL}));r.append(rp);t=el('w:t');t.text='[Enter '+tag.replace('_',' ').lower()+']';r.append(t);c.append(r);s.append(c);p._p.append(s)

class Builder:
    def __init__(self, code, spec, registry, sources):
        self.code=code;self.spec=spec;self.path=ROOT/spec['path'];self.registry=registry;self.sources=sources
        self.doc=Document();sec=self.doc.sections[0]
        for n,v in [('page_width',8.5),('page_height',11),('top_margin',.65),('bottom_margin',.65),('left_margin',.68),('right_margin',.68),('header_distance',.25),('footer_distance',.28)]:setattr(sec,n,Inches(v))
        for n in ['Normal','Body Text','List Bullet','List Number']:
            s=self.doc.styles[n];s.font.name='Calibri';s.font.size=Pt(10.5);s.font.color.rgb=RGBColor.from_string(INK);s.paragraph_format.space_after=Pt(6);s.paragraph_format.line_spacing=1.07;s.paragraph_format.widow_control=True
        for n,size,col in [('Title',25,NAVY),('Subtitle',12,TEAL),('Heading 1',17,NAVY),('Heading 2',11.5,TEAL),('Heading 3',10.5,NAVY)]:
            s=self.doc.styles[n];s.font.name='Calibri';s.font.size=Pt(size);s.font.color.rgb=RGBColor.from_string(col);s.font.bold=n!='Subtitle';s.paragraph_format.keep_with_next=True;s.paragraph_format.space_before=Pt(8);s.paragraph_format.space_after=Pt(7)
        s=self.doc.styles['Caption'];s.font.name='Calibri';s.font.size=Pt(8.5);s.font.color.rgb=RGBColor.from_string(MUTED);s.paragraph_format.space_after=Pt(5)
        sec.header.paragraphs[0].text='PORTABLE MULTI-TENANT SECURE HOSTING  |  '+code;sec.header.paragraphs[0].style='Caption'
        fp=sec.footer.paragraphs[0];fp.style='Caption';fp.text='Delivery kits v1.1  •  Reference architecture v1.4  •  Proposed   |   '
        f=el('w:fldSimple',**{'w:instr':'PAGE'});r=el('w:r');t=el('w:t');t.text='1';r.append(t);f.append(r);fp._p.append(f)
        cp=self.doc.core_properties;cp.title=spec['title'];cp.subject='Infrastructure architecture and delivery development';cp.author='';cp.comments='';cp.version='1.1'
    def p(self,text,style=None):return self.doc.add_paragraph(text,style)
    def refs(self, refs, label='Design basis and related records: '):
        p=self.p(label,'Caption')
        for i,(code,n) in enumerate(refs):
            if i:p.add_run('  •  ')
            if code in BASEMAP:
                fn,fmt=BASEMAP[code];dest=ROOT/'05_Reference_v1_4'/fn;anchor=fmt.format(n)
            else:
                dest=ROOT/(self.registry[code]['path'] if code in self.registry else OLD[code]);anchor=f'{code}_{n:02d}'
            link(p,f'{code} §{n}',os.path.relpath(dest,self.path.parent).replace(os.sep,'/'),anchor)
    def table(self,headers,rows,widths=None):
        w=widths or [7.14/len(headers)]*len(headers)
        t=self.doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
        def setcell(c,text,header=False,band=False):
            pr=c._tc.get_or_add_tcPr();mar=el('w:tcMar')
            for k,amt in [('top',85),('bottom',85),('start',90),('end',90)]:mar.append(el('w:'+k,**{'w:w':amt,'w:type':'dxa'}))
            pr.append(mar);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
            if header or band or isinstance(text,dict):pr.append(el('w:shd',**{'w:fill':NAVY if header else 'EEF5F7','w:val':'clear'}))
            p=c.paragraphs[0];p.paragraph_format.space_after=Pt(1.5);p.paragraph_format.line_spacing=1.03
            if isinstance(text,dict):response(p,text['field'])
            else:
                r=p.add_run(str(text));r.font.size=Pt(9.4);r.bold=header
                if header:r.font.color.rgb=RGBColor(255,255,255)
        for i,h in enumerate(headers):t.rows[0].cells[i].width=Inches(w[i]);setcell(t.rows[0].cells[i],h,True)
        t.rows[0]._tr.get_or_add_trPr().append(el('w:tblHeader'))
        for idx,row in enumerate(rows):
            tr=t.add_row()
            for j,txt in enumerate(row):tr.cells[j].width=Inches(w[j]);setcell(tr.cells[j],txt,band=bool(idx%2))
            tr._tr.get_or_add_trPr().append(el('w:cantSplit'))
        for col,width in zip(t.columns,w):col.width=Inches(width)
        p=self.p('');p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=.5;p.paragraph_format.line_spacing_rule=None
    def note(self,text):
        p=self.p(text);p.runs[0].bold=True;p.paragraph_format.keep_together=True;p.paragraph_format.space_before=Pt(6)
        pr=p._p.get_or_add_pPr();pr.append(el('w:shd',**{'w:fill':'EEF5F7','w:val':'clear'}))
        b=el('w:pBdr');b.append(el('w:left',**{'w:val':'single','w:sz':'12','w:space':'8','w:color':TEAL}));pr.append(b)
    def build(self):
        p=self.p('DESIGN DEVELOPMENT  /  '+self.code);p.runs[0].bold=True;p.runs[0].font.color.rgb=RGBColor.from_string(TEAL)
        self.p(self.spec['title'],'Title');self.p(self.spec['subtitle'],'Subtitle')
        self.p('Kit release v1.1 • 17 September 2026 • Parent architecture v1.4 retained')
        self.note('Proposed engineering development. Reference examples, site decisions, actual observations and approval remain separate.')
        self.p(self.spec['intro'])
        self.p('Section navigation','Heading 2')
        for n,s in enumerate(self.spec['sections'],1):
            p=self.p('');p.paragraph_format.space_after=Pt(3);link(p,f'{n}. {s["title"]}',anchor=f'{self.code}_{n:02d}')
        self.p('Basis: linked RA/WD and role-kit sections retain their original authority. The elaborations, record formats and calculations in this supplement are local proposals, not new government requirements. D-source references identify freshly checked public mechanisms, not installed compatibility. Exact source locators are in 04_Shared/development/source_reviews.csv.','Caption')
        for n,s in enumerate(self.spec['sections'],1):
            self.doc.add_page_break();p=self.p(f'{n}. {s["title"]}','Heading 1');bookmark(p,f'{self.code}_{n:02d}')
            if s.get('lead'):self.p(s['lead'])
            if s.get('refs'):self.refs(s['refs'])
            for b in s['blocks']:
                if b['kind']=='p':self.p(b['text'])
                elif b['kind']=='h':self.p(b['text'],'Heading 2')
                elif b['kind']=='note':self.note(b['text'])
                elif b['kind']=='table':self.table(b['headers'],b['rows'],b.get('widths'))
                elif b['kind']=='refs':self.refs(b['refs'],'Continue with: ')
                elif b['kind']=='external':
                    p=self.p('Verified mechanism source: ','Caption')
                    for i,sid in enumerate(b['ids']):
                        if i:p.add_run('  •  ')
                        src=self.sources[sid];link(p,sid+' — '+src['title'],src['url'])
                else:raise ValueError('Unknown block '+b['kind'])
        self.path.parent.mkdir(parents=True,exist_ok=True);self.doc.save(self.path);return self.path

def main():
    specs=json.loads((ROOT/'08_Development_Source/documents.json').read_text())
    sources={s['id']:s for s in json.loads((ROOT/'08_Development_Source/sources.json').read_text())}
    for code,spec in specs.items(): print(Builder(code,spec,specs,sources).build())
if __name__=='__main__':main()
