from pathlib import Path
import os,re,json,itertools
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from catalogue import ROOT,DOCS,BASEMAP,DATA
NAMESPACE={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
NAVY='15384A';TEAL='0B6B70';INK='21333C';MUTED='52636D';PALE='EEF5F7';LINE='C7D5DC'
bookmarks=itertools.count(1000);fields=itertools.count(7000)
def el(tag,**attrs):
 e=OxmlElement(tag)
 for k,v in attrs.items():e.set(qn(k),str(v))
 return e
def bookmark(p,name):
 i=next(bookmarks);p._p.append(el('w:bookmarkStart',**{'w:id':i,'w:name':name}));p._p.append(el('w:bookmarkEnd',**{'w:id':i}))
def hyperlink(p,text,target=None,anchor=None):
 h=el('w:hyperlink')
 if target:
  dest=target+('#'+anchor if anchor else '')
  h.set(qn('r:id'),p.part.relate_to(dest,RT.HYPERLINK,is_external=True))
 elif anchor:h.set(qn('w:anchor'),anchor)
 r=el('w:r');prop=el('w:rPr');prop.append(el('w:color',**{'w:val':TEAL}));prop.append(el('w:u',**{'w:val':'single'}));r.append(prop)
 t=el('w:t');t.text=text;r.append(t);h.append(r);p._p.append(h)
def field(p,tag):
 s=el('w:sdt');pr=el('w:sdtPr');pr.append(el('w:id',**{'w:val':next(fields)}));pr.append(el('w:tag',**{'w:val':tag}));pr.append(el('w:alias',**{'w:val':tag.replace('_',' ')}));pr.append(el('w:text',**{'w:multiLine':'1'}));s.append(pr)
 content=el('w:sdtContent');r=el('w:r');rp=el('w:rPr');rp.append(el('w:color',**{'w:val':TEAL}));r.append(rp);t=el('w:t');t.text='{{'+tag+'}}';r.append(t);content.append(r);s.append(content);p._p.append(s)
def shade(cell,color):
 pr=cell._tc.get_or_add_tcPr();pr.append(el('w:shd',**{'w:fill':color,'w:val':'clear'}))
def set_margin(cell,top=70,start=90,bottom=70,end=90):
 pr=cell._tc.get_or_add_tcPr();ms=el('w:tcMar')
 for k,v in [('top',top),('start',start),('bottom',bottom),('end',end)]:ms.append(el('w:'+k,**{'w:w':v,'w:type':'dxa'}))
 pr.append(ms)
class Book:
 def __init__(self,key):
  self.key=key;self.path=ROOT/DOCS[key][0];self.doc=Document();self.sections=[]
  sec=self.doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11);sec.top_margin=Inches(.65);sec.bottom_margin=Inches(.65);sec.left_margin=Inches(.68);sec.right_margin=Inches(.68);sec.header_distance=Inches(.26);sec.footer_distance=Inches(.28)
  st=self.doc.styles
  for n in ['Normal','Body Text','List Bullet','List Number']:
   s=st[n];s.font.name='Calibri';s.font.size=Pt(10.5);s.font.color.rgb=RGBColor.from_string(INK)
   s.paragraph_format.space_after=Pt(5);s.paragraph_format.line_spacing=1.08
  for n,size,color in [('Title',25,NAVY),('Subtitle',12,TEAL),('Heading 1',17,NAVY),('Heading 2',12,TEAL),('Heading 3',10.5,NAVY)]:
   s=st[n];s.font.name='Calibri';s.font.size=Pt(size);s.font.color.rgb=RGBColor.from_string(color);s.font.bold=n!='Subtitle';s.paragraph_format.space_before=Pt(8);s.paragraph_format.space_after=Pt(7);s.paragraph_format.keep_with_next=True
  for n in ['Caption']:
   st[n].font.name='Calibri';st[n].font.size=Pt(8.5);st[n].font.color.rgb=RGBColor.from_string(MUTED)
  hp=sec.header.paragraphs[0];hp.text='PORTABLE MULTI-TENANT SECURE HOSTING  |  '+key;hp.style='Caption'
  fp=sec.footer.paragraphs[0];fp.text='Delivery kits v1.0  •  Architecture baseline v1.4  •  Proposed   |   ';fp.style='Caption'
  fld=el('w:fldSimple',**{'w:instr':'PAGE'});rr=el('w:r');tt=el('w:t');tt.text='1';rr.append(tt);fld.append(rr);fp._p.append(fld)
  cp=self.doc.core_properties;cp.title=DOCS[key][1];cp.subject='Infrastructure architecture, engineering and implementation delivery kit';cp.author='';cp.keywords='Portable hosting, infrastructure, architecture, engineering, implementation';cp.comments=''
 def p(self,text='',boldstart=None,style=None):
  p=self.doc.add_paragraph(style=style)
  if boldstart and text.startswith(boldstart):p.add_run(boldstart).bold=True;p.add_run(text[len(boldstart):])
  else:p.add_run(text)
  return p
 def h(self,text):return self.doc.add_paragraph(text,'Heading 2')
 def note(self,text):
  p=self.p(text);p.runs[0].bold=True
  p.paragraph_format.space_before=Pt(7);p.paragraph_format.space_after=Pt(10)
  p.paragraph_format.keep_together=True
  pr=p._p.get_or_add_pPr();pr.append(el('w:shd',**{'w:fill':PALE,'w:val':'clear'}))
  borders=el('w:pBdr');borders.append(el('w:left',**{'w:val':'single','w:sz':'12','w:space':'8','w:color':TEAL}));pr.append(borders)
 def table(self,headers,rows,widths=None):
  t=self.doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
  w=widths or [7.14/len(headers)]*len(headers)
  for i,h in enumerate(headers):
   c=t.rows[0].cells[i];c.width=Inches(w[i]);c.text=h;shade(c,NAVY);set_margin(c)
   for r in c.paragraphs[0].runs:r.bold=True;r.font.color.rgb=RGBColor(255,255,255);r.font.size=Pt(9)
  t.rows[0]._tr.get_or_add_trPr().append(el('w:tblHeader'))
  for ri,row in enumerate(rows):
   cells=t.add_row().cells
   for j,val in enumerate(row):
    c=cells[j];c.width=Inches(w[j]);set_margin(c);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
    p=c.paragraphs[0];p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=1.02
    if isinstance(val,dict) and 'field' in val:shade(c,'F0F7F7');field(p,val['field'])
    else:p.add_run(str(val))
    for r in p.runs:r.font.size=Pt(9.3)
    if ri%2==1 and not isinstance(val,dict):shade(c,'F4F7F9')
   t.rows[-1]._tr.get_or_add_trPr().append(el('w:cantSplit'))
  # apply fixed grid for reliable rendering
  for col,ww in zip(t.columns,w):col.width=Inches(ww)
  self.p().paragraph_format.space_after=Pt(0)
  return t
 def inputs(self,rows):
  return self.table(['Record / decision','What to enter','Working response'],[(a,b,{'field':c}) for a,b,c in rows],[1.53,3.70,1.91])
 def refs(self,refs):
  p=self.p('Baseline and related records: ',style='Caption')
  for i,(code,section) in enumerate(refs):
   if i:p.add_run('  •  ')
   if code in BASEMAP:
    fn,fmt=BASEMAP[code];target=ROOT/'05_Reference_v1_4'/fn;anchor=fmt.format(section)
   else:target=ROOT/DOCS[code][0];anchor=f'{code}_{section:02d}'
   hyperlink(p,f'{code} §{section}',os.path.relpath(target,self.path.parent).replace(os.sep,'/'),anchor)
 def external(self,ids):
  p=self.p('External mechanism context: ',style='Caption')
  for i,sid in enumerate(ids):
   if i:p.add_run('  •  ')
   src=next(x for x in DATA['sources'] if x['id']==sid);hyperlink(p,sid+' — '+src['title'],src['url'])
 def page(self,num,title,lead=None,refs=None):
  if self.sections:self.doc.add_page_break()
  p=self.doc.add_paragraph(f'{num}. {title}','Heading 1');bookmark(p,f'{self.key}_{num:02d}');self.sections.append((num,title))
  if lead:self.p(lead)
  if refs:self.refs(refs)
 def front(self,subtitle,how,contents):
  p=self.p('DELIVERY KIT  /  '+self.key);p.runs[0].font.color.rgb=RGBColor.from_string(TEAL);p.runs[0].bold=True
  self.doc.add_paragraph(DOCS[self.key][1],'Title');self.doc.add_paragraph(subtitle,'Subtitle')
  self.p('Kit v1.0 • 16 September 2026 • Aligned to the frozen v1.4 reference architecture')
  self.note('Working kit, not an approved site design or deployed platform. Templates, reference examples and live evidence are separate records.')
  self.p(how)
  self.h('Section links')
  for num,title in contents:
   p=self.p();p.paragraph_format.space_after=Pt(3);hyperlink(p,f'{num}. {title}',anchor=f'{self.key}_{num:02d}')
  self.p('All additional process guidance is proposed kit practice. RA/WD references identify baseline-derived architecture; K references identify external mechanism checks. No site values or approval signatures are supplied.',style='Caption')
  self.sections=[(0,'Front')]
 def save(self):
  self.path.parent.mkdir(parents=True,exist_ok=True);self.doc.save(self.path)
  return str(self.path)
